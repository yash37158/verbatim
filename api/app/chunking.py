"""Parse a document into pages, then into overlapping chunks with exact offsets.

The offsets are the point. A citation resolves to a page and a character span in the
document's own text, so every transformation here stays reversible: for every chunk,
`full_text[chunk.char_start:chunk.char_end] == chunk.text`.
"""

import re
from bisect import bisect_right
from dataclasses import dataclass

from .config import settings

# Coarsest first: prefer to cut between paragraphs, then lines, then sentences, then words.
SEPARATORS = ["\n\n", "\n", ". ", " "]

# A text-layer PDF averages thousands of characters per page. A scan averages a handful.
MIN_CHARS_PER_PAGE = 100


class UnreadableDocument(Exception):
    """Raised when there is no text to index — surfaced to the user, not retried."""


@dataclass(frozen=True)
class Page:
    number: int | None  # None for formats without real pagination (txt, md, docx)
    text: str


@dataclass(frozen=True)
class Chunk:
    ordinal: int
    text: str
    char_start: int
    char_end: int
    page_start: int | None
    page_end: int | None


def extract(data: bytes, filename: str, mime_type: str = "") -> list[Page]:
    name = filename.lower()
    if name.endswith(".pdf") or mime_type == "application/pdf":
        return _extract_pdf(data)
    if name.endswith(".docx"):
        return _extract_docx(data)
    if name.endswith((".txt", ".md", ".markdown")):
        return [Page(None, data.decode("utf-8", errors="replace"))]
    raise UnreadableDocument(f"Unsupported file type: {filename}")


def _extract_pdf(data: bytes) -> list[Page]:
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as doc:
        pages = [Page(i + 1, page.get_text()) for i, page in enumerate(doc)]

    if not pages:
        raise UnreadableDocument("This PDF has no pages.")
    if sum(len(p.text.strip()) for p in pages) / len(pages) < MIN_CHARS_PER_PAGE:
        raise UnreadableDocument(
            "No extractable text layer — this looks like a scan. OCR is not available yet."
        )
    return pages


def _extract_docx(data: bytes) -> list[Page]:
    import io

    import docx

    doc = docx.Document(io.BytesIO(data))
    text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if not text.strip():
        raise UnreadableDocument("This document contains no text.")
    # python-docx has no concept of a page — pagination is decided by the renderer.
    return [Page(None, text)]


def chunk_pages(
    pages: list[Page],
    limit: int | None = None,
    overlap: int | None = None,
    max_pages: int | None = None,
) -> tuple[str, list[Chunk]]:
    """Return the flattened document text and its chunks."""
    limit = limit or settings.chunk_chars
    overlap = settings.chunk_overlap_chars if overlap is None else overlap
    max_pages = max_pages or settings.chunk_max_pages

    full, bounds = _flatten(pages)
    starts = [b[0] for b in bounds]

    def page_at(pos: int) -> int | None:
        i = max(0, bisect_right(starts, pos) - 1)
        return bounds[i][2]

    chunks: list[Chunk] = []
    spans = (
        span
        for group_start, group_end in _page_groups(bounds, limit, max_pages)
        for span in _spans(full, group_start, group_end, limit, SEPARATORS)
    )
    for start, end in spans:
        begin = _extend_back(full, start, overlap)
        text = full[begin:end]
        if not text.strip():
            continue
        chunks.append(
            Chunk(
                ordinal=len(chunks),
                text=text,
                char_start=begin,
                char_end=end,
                # Pages come from the core span, not `begin`: the overlap reaches back
                # into the previous page as context, and a citation must name the page
                # the content is actually on.
                page_start=page_at(start),
                page_end=page_at(max(start, end - 1)),
            )
        )
    return full, chunks


def page_offsets(pages: list[Page]) -> list[tuple[int, int | None]]:
    """Where each page starts in the flattened text — how a citation resolves a quote's
    offset back to a page number."""
    return [(start, number) for start, _, number in _flatten(pages)[1]]


def _flatten(pages: list[Page]) -> tuple[str, list[tuple[int, int, int | None]]]:
    parts: list[str] = []
    bounds: list[tuple[int, int, int | None]] = []
    pos = 0
    for i, page in enumerate(pages):
        if i:
            parts.append("\n\n")
            pos += 2
        parts.append(page.text)
        bounds.append((pos, pos + len(page.text), page.number))
        pos += len(page.text)
    return "".join(parts), bounds


def _page_groups(
    bounds: list[tuple[int, int, int | None]], limit: int, max_pages: int
) -> list[tuple[int, int]]:
    """Batch consecutive pages so no chunk can range further than a citation can name.

    On a dense document every page already exceeds `limit`, so each group is one page and
    this changes nothing. It only bites on sparse pages, where the character limit alone
    would happily swallow the whole file into a single chunk.
    """
    if not bounds:
        return []
    groups: list[tuple[int, int]] = []
    start, pages = bounds[0][0], 0
    for i, (_, page_end, _page_no) in enumerate(bounds):
        pages += 1
        if page_end - start >= limit or pages >= max_pages:
            groups.append((start, page_end))
            # Resume at the next page's own offset, not this one's end — the separator
            # between them belongs to neither, and would misattribute the next chunk.
            start = bounds[i + 1][0] if i + 1 < len(bounds) else page_end
            pages = 0
    if start < bounds[-1][1]:
        groups.append((start, bounds[-1][1]))
    return groups


def _cut_points(text: str, start: int, end: int, sep: str) -> list[int]:
    """Offsets just past each separator occurrence, strictly inside (start, end)."""
    out: list[int] = []
    p = text.find(sep, start, end)
    while p != -1:
        cut = p + len(sep)
        if start < cut < end:
            out.append(cut)
        p = text.find(sep, cut, end)
    return out


def _spans(text: str, start: int, end: int, limit: int, seps: list[str]):
    """Contiguous (start, end) spans covering [start, end), each at most `limit` chars."""
    if end - start <= limit:
        if end > start:
            yield (start, end)
        return

    for i, sep in enumerate(seps):
        cuts = _cut_points(text, start, end, sep)
        if not cuts:
            continue
        rest = seps[i + 1 :]
        piece = prev = start
        for cut in (*cuts, end):
            if cut - piece > limit:
                if prev > piece:  # flush the runs that did fit
                    yield (piece, prev)
                    piece = prev
                if cut - piece > limit:  # a single run is over the limit on its own
                    yield from _spans(text, piece, cut, limit, rest)
                    piece = cut
            prev = cut
        if piece < end:
            yield from _spans(text, piece, end, limit, rest)
        return

    for s in range(start, end, limit):  # no separator helped
        yield (s, min(s + limit, end))


def _extend_back(text: str, start: int, overlap: int) -> int:
    """Pull the chunk's start back by `overlap`, snapped forward so it never begins mid-word."""
    if start == 0 or overlap <= 0:
        return start
    s = max(0, start - overlap)
    m = re.search(r"\s", text[s:start])
    return s + m.end() if m else s
